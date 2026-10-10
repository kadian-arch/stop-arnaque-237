"""Release formats besides JSONL/CSV: a readable Excel workbook and Parquet (what Hugging Face serves)."""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .schema import DOC, V

RED, NAVY = "D32F2F", "0B132B"
WIDE = {"text", "message_text", "call_description", "extra_notes", "summary", "example_message"}
MEDIUM = {"title", "scam_types", "top_scam_types", "source_url", "sender", "domains", "phone_ids", "entities",
          "caller_asked", "actions_after", "requested_actions", "impersonated", "sender_shown", "meaning", "allowed values"}
ROW_HEIGHT = 48   # points: about three lines; longer text stays inside the cell (full text in the formula bar)


def _cell(v):
    if isinstance(v, list):
        return " | ".join(map(str, v))
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    return v


def _sheet(wb, name, fields, rows):
    ws = wb.create_sheet(name)
    ws.append(fields)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append([_cell(r.get(f)) for f in fields])
    for i, f in enumerate(fields, 1):
        ws.column_dimensions[get_column_letter(i)].width = 80 if f in WIDE else 32 if f in MEDIUM else max(12, min(22, len(f) + 4))
    ws.row_dimensions[1].height = 22
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = ROW_HEIGHT
        for c, f in zip(row, fields):
            c.alignment = Alignment(vertical="top", wrap_text=f in WIDE or f in MEDIUM)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def workbook(path: Path, version: str, tables: dict, fields: dict, stats: dict):
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    lines = [
        ("Stop Arnaque 237: Cameroon Scam Reports Dataset", True),
        (f"Version {version}", False),
        ("", False),
        ("Scam and genuine messages from Cameroon, labelled. Start with the messages sheet; the other sheets hold the full records behind it.", False),
        ("Everything is anonymized: phone numbers are replaced by keyed codes, names, account numbers and ids are removed, links are made unclickable.", False),
        ("", False),
        ("Sheets", True),
    ] + [(f"{name}: {len(rows)} rows", False) for name, rows in tables.items()] + [
        ("data_dictionary: what each column means and the values it can take", False),
        ("", False),
        ("Licence: CC BY 4.0. Credit: Stop Arnaque 237: Cameroon Scam Reports Dataset, https://github.com/kadian-arch/stop-arnaque-237", False),
        ("Contact and removal requests: groundtruth.cm@gmail.com", False),
        ("Full documentation: DATASHEET.md in the repository.", False),
    ]
    for text, bold in lines:
        ws.append([text])
        if bold:
            ws.cell(ws.max_row, 1).font = Font(bold=True, size=14 if ws.max_row == 1 else 12, color=RED if ws.max_row == 1 else NAVY)
    ws.column_dimensions["A"].width = 130

    for name, rows in tables.items():
        _sheet(wb, name, fields[name], rows)

    dd = wb.create_sheet("data_dictionary")
    dd.append(["table", "column", "meaning", "allowed values"])
    for c in dd[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
    for name in tables:
        for f in fields[name]:
            dd.append([name, f, DOC.get(name, {}).get(f, ""), ", ".join(V[f]) if f in V else ""])
    for col, w in zip("ABCD", (18, 22, 90, 60)):
        dd.column_dimensions[col].width = w
    for row in dd.iter_rows(min_row=2):
        dd.row_dimensions[row[0].row].height = ROW_HEIGHT
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    dd.freeze_panes = "A2"
    wb.save(path)


def parquet(path: Path, fields, rows):
    import pyarrow as pa
    import pyarrow.parquet as pq
    cols = {f: [r.get(f) for r in rows] for f in fields}
    # keep lists as lists, everything else as text/number; mixed columns fall back to JSON text
    arrays = {}
    for f, vals in cols.items():
        try:
            arrays[f] = pa.array(vals)
        except (pa.ArrowInvalid, pa.ArrowTypeError):
            arrays[f] = pa.array([None if v is None else json.dumps(v, ensure_ascii=False) for v in vals])
    pq.write_table(pa.table(arrays), path)
