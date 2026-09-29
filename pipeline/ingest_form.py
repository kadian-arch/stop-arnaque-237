"""Tally CSV export -> normalized report records (still raw, not yet anonymized).

Columns are found by the English start of each question title, so the importer
survives wording tweaks. Multi-choice answers are read in either layout Tally may use:
one column holding all ticked options, or one column per option ("Title (Option)") with true/false.
"""
import csv
import hashlib
import json
import re
from pathlib import Path

from . import taxonomy as T

# key -> (title prefix, kind, mapping table)
QUESTIONS = {
    "age_ok": ("i am 18 or older", "flag", None),
    "consent": ("i agree that my report", "flag", None),
    "channel": ("how did it reach you", "one", T.FORM_CHANNEL),
    "scam_types": ("what was it about", "many", T.FORM_SCAM_TYPE),
    "screenshots": ("upload screenshots", "files", None),
    "message_text": ("or paste the message text", "text", None),
    "sender": ("who sent it", "text", None),
    "sender_type": ("the sender looked like", "one", T.FORM_SENDER_TYPE),
    "caller_claimed": ("who did the caller say", "one", T.FORM_CALLER_CLAIM),
    "caller_asked": ("what did they ask you to do", "many", T.FORM_REQUESTED),
    "call_language": ("language of the call", "one", T.FORM_CALL_LANGUAGE),
    "call_end": ("how did it end", "one", T.FORM_CALL_END),
    "caller_number": ("caller's number", "text", None),
    "call_description": ("describe the call", "text", None),
    "outcome": ("did you lose money", "one", T.FORM_OUTCOME),
    "amount_lost_fcfa": ("how much, in fcfa", "number", None),
    "payment_rails": ("through what", "many", T.FORM_RAIL),
    "actions_after": ("what did you do after", "many", T.FORM_AFTER),
    "when": ("when did it happen", "one", T.FORM_WHEN),
    "region": ("your region", "one", T.FORM_REGION),
    "extra_notes": ("anything else we should know", "text", None),
    "would_use_tool": ("would you use a free tool", "one", T.FORM_WOULD_USE),
    "wants_report_email": ("want our free public report", "text", None),
    "src": ("src", "text", None),
}

TRUE = {"true", "yes", "1", "x", "checked", "oui", "vrai"}


def _norm(h: str) -> str:
    return re.sub(r"\s+", " ", (h or "").strip().lower().replace("’", "'"))


def _columns(headers):
    """question key -> list of (header, option-or-None)."""
    cols = {k: [] for k in QUESTIONS}
    for h in headers:
        n = _norm(h)
        # Google Sheets integration: a checkbox block with no title comes out as
        # "Untitled checkboxes field (<option>)", so the option is the question
        u = re.match(r"untitled [a-z ]*field \((.*)\)$", n)
        if u:
            for key, (prefix, kind, _) in QUESTIONS.items():
                if kind == "flag" and u.group(1).startswith(prefix):
                    cols[key].append((h, None))
                    break
            continue
        for key, (prefix, _, _) in QUESTIONS.items():
            if key == "src":
                if n == "src":
                    cols[key].append((h, None))
                continue
            if n.startswith(prefix):
                m = re.search(r"\(([^()]*(?:\([^()]*\)[^()]*)*)\)\s*$", h or "")
                opt = m.group(1) if m and len(h) - len(m.group(0)) >= len(prefix) else None
                cols[key].append((h, opt))
                break
    return cols


def _value(row, key, cols):
    prefix, kind, table = QUESTIONS[key]
    found = cols.get(key) or []
    if not found:
        return None
    if kind in ("many", "one") and any(opt for _, opt in found):
        ticked = [opt for h, opt in found if opt and str(row.get(h, "")).strip().lower() in TRUE]
        codes = []
        for opt in ticked:
            c = T.match(opt, table)
            if c and c not in codes:
                codes.append(c)
        if kind == "one":
            return codes[0] if codes else None
        return codes
    cell = str(row.get(found[0][0], "") or "").strip()
    if kind == "flag":
        return cell.lower() in TRUE or (cell != "" and cell.lower() not in {"false", "no", "0"})
    if kind == "one":
        return T.match(cell, table)
    if kind == "many":
        return T.match_many(cell, table)
    if kind == "number":
        return parse_amount(cell)
    if kind == "files":
        return re.findall(r"https?://[^\s,()\"']+", cell)
    return cell or None


def parse_amount(cell):
    """'25000', '25 000', '25.000', '25000.00', '25k', '2,5 millions', '1.5m' -> int FCFA."""
    s = str(cell or "").strip().lower().replace(" ", " ")
    if not s:
        return None
    mult = 1
    if re.search(r"\d\s*(k|mille)\b", s):
        mult = 1_000
    elif re.search(r"\d\s*(m|mil|million|millions|mio)\b", s):
        mult = 1_000_000
    num = re.search(r"\d[\d\s.,']*", s)
    if not num:
        return None
    t = num.group(0).strip().replace(" ", "").replace("'", "")
    if re.fullmatch(r"\d{1,3}([.,]\d{3})+", t):          # thousands separators: 25.000 / 1,500,000
        t = re.sub(r"[.,]", "", t)
    elif re.fullmatch(r"\d+[.,]\d{1,2}", t):             # decimals: 25000.00 / 2,5
        t = t.replace(",", ".")
    else:
        t = re.sub(r"[.,]", "", t)
    try:
        return int(round(float(t) * mult))
    except ValueError:
        return None


def load_table(path: Path):
    """(headers, rows) from a Tally .csv (comma or semicolon, UTF-8 or Windows encoding,
    e.g. after being opened and re-saved in Excel) or an .xlsx export."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        headers = [str(h or "").strip() for h in next(it)]
        rows = [{h: ("" if v is None else str(v)) for h, v in zip(headers, r)} for r in it if any(v not in (None, "") for v in r)]
        return headers, rows
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    first = text.splitlines()[0] if text else ""
    delim = ";" if first.count(";") > first.count(",") else ","
    reader = csv.DictReader(text.splitlines(keepends=True), delimiter=delim)
    headers = reader.fieldnames or []
    rows = [r for r in reader if any((v or "").strip() for v in r.values() if isinstance(v, str))]
    return headers, rows


def read_tally_csv(path: Path) -> list:
    headers, rows = load_table(Path(path))
    cols = _columns(headers)
    out = []
    for row in rows:
        sid = row.get("Submission ID") or row.get("Response ID") or row.get("ID")
        if not sid:  # never seen in Tally exports, but never lose a row over it
            sid = "nosid-" + hashlib.sha1(json.dumps(row, sort_keys=True).encode()).hexdigest()[:10]
        rec = {
            "record_type": "report",
            "source": "form",
            "submission_id": sid,
            "submitted_at": row.get("Submitted at") or row.get("Created at") or row.get("Date"),
        }
        for key in QUESTIONS:
            rec[key] = _value(row, key, cols)
        rec["src"] = rec.get("src") or "direct"
        out.append(rec)
    return out


def missing_columns(path: Path) -> list:
    """Questions the importer could not find in this export (checked on every run)."""
    headers, _ = load_table(Path(path))
    cols = _columns(headers)
    return [k for k, v in cols.items() if not v and k != "src"]
